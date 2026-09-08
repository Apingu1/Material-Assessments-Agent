from __future__ import annotations

import math
import re
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .coshh_rules import AREA_LABELS
from .models import COSHHActivityAssessment, COSHHArea, COSHHAssessment, GHSCode, MaterialInput, RiskBand
from .sds import StoredSDS


NAVY = "17365D"
LIGHT_TEAL = "DCEFF2"
PALE = "F5F8FA"
MID_GREY = "6B7280"
DARK = "1F2937"
RED = "C00000"
AMBER = "F4B183"
GREEN = "A9D18E"
YELLOW = "FFE699"
WHITE = "FFFFFF"
BORDER = "D4DCE3"


def _shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _border(cell, color: str = BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "bottom", "left", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)
        borders.append(node)


def _margins(cell, value: int = 60) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, width in (("top", value), ("start", 80), ("bottom", value), ("end", 80)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(width))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def _text(cell, value: str, *, bold: bool = False, size: float = 8.1, color: str = DARK, align=None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(value or "")
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def _para(target, value: str, *, bold: bool = False, size: float = 8.1, color: str = DARK, after: float = 1, align=None):
    paragraph = target.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(value or "")
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def _style_table(table, widths: list[float], margin: int = 55) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _margins(cell, margin)
            _border(cell)


def _repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


def _no_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _risk_letter(value: RiskBand) -> str:
    return {RiskBand.LOW: "L", RiskBand.MEDIUM: "M", RiskBand.HIGH: "H"}[value]


def _risk_fill(value: RiskBand) -> str:
    return {RiskBand.LOW: GREEN, RiskBand.MEDIUM: YELLOW, RiskBand.HIGH: AMBER}[value]


def _clean_lines(values: list[str], limit: int = 3) -> str:
    cleaned = [" ".join(value.split()).strip(" .") for value in values if value.strip()]
    return "; ".join(cleaned[:limit]) + ("." if cleaned else "")


def _hazard_line(values: list[str], limit: int = 5) -> str:
    values = [" ".join(value.split()) for value in values if value.strip()]
    shown = values[:limit]
    suffix = f" (+{len(values) - limit} more in research record)" if len(values) > limit else ""
    return "\n".join(shown) + suffix


def _draw_diamond(draw: ImageDraw.ImageDraw) -> None:
    diamond = [(250, 25), (475, 250), (250, 475), (25, 250)]
    draw.polygon(diamond, fill="white", outline=(210, 0, 0), width=28)


def _draw_flame(draw: ImageDraw.ImageDraw, x: int = 250, y: int = 275, scale: float = 1.0) -> None:
    pts = [
        (x, y - int(145 * scale)),
        (x + int(58 * scale), y - int(55 * scale)),
        (x + int(40 * scale), y + int(5 * scale)),
        (x + int(92 * scale), y + int(58 * scale)),
        (x + int(32 * scale), y + int(130 * scale)),
        (x - int(30 * scale), y + int(135 * scale)),
        (x - int(88 * scale), y + int(75 * scale)),
        (x - int(48 * scale), y + int(8 * scale)),
        (x - int(65 * scale), y - int(62 * scale)),
    ]
    draw.polygon(pts, fill="black")
    draw.polygon([(x, y - 55), (x + 28, y + 20), (x, y + 82), (x - 28, y + 22)], fill="white")


def _make_ghs_icon(code: GHSCode, path: Path) -> None:
    image = Image.new("RGBA", (500, 500), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    _draw_diamond(draw)

    if code == GHSCode.GHS01:
        draw.ellipse((190, 210, 300, 320), fill="black")
        for angle in range(0, 360, 40):
            a = math.radians(angle)
            x1, y1 = 245 + 70 * math.cos(a), 265 + 70 * math.sin(a)
            x2, y2 = 245 + 155 * math.cos(a), 265 + 155 * math.sin(a)
            draw.line((x1, y1, x2, y2), fill="black", width=18)
    elif code == GHSCode.GHS02:
        _draw_flame(draw)
    elif code == GHSCode.GHS03:
        draw.ellipse((185, 300, 315, 430), outline="black", width=24)
        _draw_flame(draw, 250, 225, 0.62)
    elif code == GHSCode.GHS04:
        draw.rounded_rectangle((120, 220, 370, 300), radius=30, outline="black", width=24)
        draw.rectangle((355, 238, 405, 282), fill="black")
    elif code == GHSCode.GHS05:
        draw.line((100, 360, 400, 360), fill="black", width=18)
        draw.polygon([(160, 315), (300, 315), (270, 355), (140, 355)], fill="black")
        draw.line((135, 180, 215, 265), fill="black", width=20)
        draw.line((305, 180, 260, 275), fill="black", width=20)
        draw.ellipse((205, 270, 235, 300), fill="black")
        draw.ellipse((245, 285, 275, 315), fill="black")
    elif code == GHSCode.GHS06:
        draw.ellipse((180, 135, 320, 275), fill="black")
        draw.ellipse((210, 170, 245, 205), fill="white")
        draw.ellipse((255, 170, 290, 205), fill="white")
        draw.polygon([(235, 210), (265, 210), (250, 235)], fill="white")
        draw.line((150, 335, 350, 405), fill="black", width=25)
        draw.line((350, 335, 150, 405), fill="black", width=25)
    elif code == GHSCode.GHS07:
        draw.rounded_rectangle((226, 125, 274, 315), radius=20, fill="black")
        draw.ellipse((223, 350, 277, 404), fill="black")
    elif code == GHSCode.GHS08:
        draw.ellipse((205, 105, 295, 195), fill="black")
        draw.polygon([(160, 385), (185, 240), (225, 210), (275, 210), (315, 240), (340, 385)], fill="black")
        draw.ellipse((213, 235, 287, 315), fill="white")
        points = []
        for index in range(10):
            angle = -math.pi / 2 + index * math.pi / 5
            radius = 34 if index % 2 == 0 else 14
            points.append((250 + radius * math.cos(angle), 275 + radius * math.sin(angle)))
        draw.polygon(points, fill="black")
    elif code == GHSCode.GHS09:
        draw.line((160, 360, 220, 155), fill="black", width=22)
        draw.line((220, 155, 250, 360), fill="black", width=16)
        draw.line((205, 215, 155, 180), fill="black", width=12)
        draw.line((210, 250, 270, 205), fill="black", width=12)
        draw.line((130, 365, 380, 365), fill="black", width=10)
        draw.ellipse((300, 300, 360, 335), outline="black", width=10)
        draw.polygon([(360, 318), (398, 293), (390, 342)], fill="black")

    image.save(path)


def _icon(code: GHSCode, icon_dir: Path) -> Path:
    icon_dir.mkdir(parents=True, exist_ok=True)
    path = icon_dir / f"{code.value}.png"
    if not path.exists():
        _make_ghs_icon(code, path)
    return path


def _compact_activities(assessment: COSHHAssessment) -> list[COSHHActivityAssessment]:
    """Keep the signed COSHH concise while the JSON retains the full task-level assessment."""
    activities = assessment.activities
    result: list[COSHHActivityAssessment] = []

    def first(area: COSHHArea, contains: str) -> COSHHActivityAssessment | None:
        contains = contains.lower()
        return next(
            (x for x in activities if x.area == area and contains in x.activity.lower()),
            None,
        )

    for area in (COSHHArea.GOODS_IN_WAREHOUSE,):
        for term in ("check outer", "open outer"):
            found = first(area, term)
            if found:
                result.append(found)

    preferred = {
        COSHHArea.SAMPLING: "open / sample",
        COSHHArea.QC: "weigh / prepare",
        COSHHArea.PRODUCTION: "weighing",
        COSHHArea.R_AND_D: "weighing",
        COSHHArea.HOUSEKEEPING: "spillage",
        COSHHArea.OFFICE: "administrative",
    }
    for area, term in preferred.items():
        found = first(area, term)
        if found:
            result.append(found)

    # Show one universal spill-response row when relevant and not already represented by Housekeeping.
    spill = next((x for x in activities if "spill" in x.activity.lower()), None)
    if spill and not any("spill" in x.activity.lower() for x in result):
        result.append(spill)
    return result


def build_coshh_docx(
    assessment: COSHHAssessment,
    item: MaterialInput,
    stored_sds: StoredSDS,
    output_path: Path,
) -> None:
    research = assessment.research
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.38)
    section.bottom_margin = Inches(0.32)
    section.left_margin = Inches(0.46)
    section.right_margin = Inches(0.46)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(8.7)

    if "SectionTitle" not in doc.styles:
        style = doc.styles.add_style("SectionTitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["SectionTitle"]
    style.font.name = "Arial"
    style.font.size = Pt(10.8)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(NAVY)
    style.paragraph_format.space_before = Pt(4)
    style.paragraph_format.space_after = Pt(2)

    header = doc.add_table(rows=1, cols=3)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    header.autofit = False
    for index, width in enumerate((1.65, 4.55, 1.25)):
        header.cell(0, index).width = Inches(width)
    for cell in header.rows[0].cells:
        _shading(cell, NAVY)
        _margins(cell, 85)
    _text(header.cell(0, 0), "EASTSTONE", bold=True, size=14, color=WHITE)
    _text(header.cell(0, 1), "COSHH ASSESSMENT", bold=True, size=16, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    _text(header.cell(0, 2), "DRAFT", bold=True, size=9, color=WHITE, align=WD_ALIGN_PARAGRAPH.RIGHT)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run("AGENT-GENERATED DRAFT  |  HUMAN REVIEW & SIGNATURE REQUIRED")
    run.font.name = "Arial"
    run.font.size = Pt(8.8)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(RED)

    doc.add_paragraph("1. Substance & Assessment Scope", style="SectionTitle")
    identity = doc.add_table(rows=5, cols=4)
    _style_table(identity, [1.25, 2.45, 1.25, 2.45])
    rows = [
        ("Substance", research.coshh_substance_name, "Material / use context", research.material_use_context),
        ("CAS No.", research.sds.cas_number or "Not confirmed", "Assessment type", "Routine COSHH - draft"),
        ("SDS used", research.sds.manufacturer or "Reference SDS", "SDS revision", research.sds.revision_date or "Not stated"),
        ("Applicable areas", ", ".join(AREA_LABELS[a] for a in item.coshh_areas), "Frequency / duration", f"{item.frequency} / {item.duration}"),
        ("People exposed", item.people_exposed or "To be confirmed", "Typical quantity", item.typical_quantity or "To be confirmed"),
    ]
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = identity.cell(row_index, col_index)
            if col_index % 2 == 0:
                _shading(cell, LIGHT_TEAL)
                _text(cell, value, bold=True, size=7.9, color=NAVY)
            else:
                _text(cell, value, size=8.1)

    doc.add_paragraph("2. Hazard Summary", style="SectionTitle")
    hazards = doc.add_table(rows=1, cols=3)
    _style_table(hazards, [1.75, 2.25, 3.45], margin=50)
    icon_cell = hazards.cell(0, 0)
    _shading(icon_cell, PALE)
    icon_paragraph = icon_cell.paragraphs[0]
    icon_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    icon_dir = output_path.parent / ".coshh_icons"
    for index, code in enumerate(research.pictograms[:4]):
        if index:
            icon_paragraph.add_run(" ")
        icon_paragraph.add_run().add_picture(str(_icon(code, icon_dir)), width=Inches(0.55))
    if not research.pictograms:
        _text(icon_cell, "No pictogram identified", size=8, color=MID_GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

    class_cell = hazards.cell(0, 1)
    _shading(class_cell, "FFF7F7")
    _text(class_cell, f"Signal word: {research.signal_word}", bold=True, size=8.7, color=RED if research.signal_word in {"DANGER", "WARNING"} else DARK)
    _para(class_cell, _hazard_line(research.hazard_statements, 5) or "No H-statements identified", size=7.8, after=0)

    summary_cell = hazards.cell(0, 2)
    _text(summary_cell, "What matters for Eaststone", bold=True, size=8.7, color=NAVY)
    for control in research.core_controls[:3]:
        _para(summary_cell, f"• {control}", size=7.9)
    if research.workplace_exposure_limit_found:
        _para(summary_cell, f"• WEL identified: {research.workplace_exposure_limit}", size=7.9, color=RED)

    doc.add_paragraph("3. Exposure Routes", style="SectionTitle")
    exposure = doc.add_table(rows=5, cols=3)
    _style_table(exposure, [1.45, 1.05, 4.9], margin=50)
    for index, value in enumerate(("Route", "Relevant?", "Assessment")):
        _shading(exposure.cell(0, index), NAVY)
        _text(exposure.cell(0, index), value, bold=True, size=7.9, color=WHITE)
    route_rows = [
        ("Inhalation", research.inhalation_hazard, "Intrinsic inhalation hazard and/or credible airborne exposure during open handling."),
        ("Skin", research.skin_hazard, "Skin exposure is considered for direct handling, cleaning and spill response."),
        ("Eyes", research.eye_hazard, "Eye exposure is considered where dust/splash is foreseeable."),
        ("Ingestion", research.ingestion_hazard, "Not a normal workplace route; included where the SDS identifies ingestion hazard."),
    ]
    for row_index, (route, relevant, note) in enumerate(route_rows, 1):
        _text(exposure.cell(row_index, 0), route, bold=True, size=7.9, color=NAVY)
        _text(exposure.cell(row_index, 1), "YES" if relevant else "NO", bold=True, size=7.9, color=RED if relevant else DARK, align=WD_ALIGN_PARAGRAPH.CENTER)
        _text(exposure.cell(row_index, 2), note, size=7.9)

    doc.add_paragraph("4. Core Controls", style="SectionTitle")
    controls_table = doc.add_table(rows=4, cols=2)
    _style_table(controls_table, [1.55, 5.85], margin=50)
    controls = [
        ("Containment", "Keep containers closed when not in use; minimise the amount exposed during handling."),
        ("Engineering", item.existing_controls or "Confirm the applicable local containment / extraction arrangement before approval."),
        ("PPE", "Use task-specific PPE shown in Section 5; RPE is supplementary to adequate engineering controls."),
        ("Hygiene", "No eating/drinking in the work area; wash hands after handling and before breaks."),
    ]
    for index, (label, value) in enumerate(controls):
        _shading(controls_table.cell(index, 0), LIGHT_TEAL)
        _text(controls_table.cell(index, 0), label, bold=True, size=7.9, color=NAVY)
        _text(controls_table.cell(index, 1), value, size=7.9)

    doc.add_page_break()

    page_header = doc.add_table(rows=1, cols=2)
    page_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    page_header.autofit = False
    page_header.cell(0, 0).width = Inches(5.7)
    page_header.cell(0, 1).width = Inches(1.7)
    for cell in page_header.rows[0].cells:
        _shading(cell, NAVY)
        _margins(cell, 70)
    _text(page_header.cell(0, 0), "COSHH ACTIVITY & CONTROL ASSESSMENT", bold=True, size=12.5, color=WHITE)
    _text(page_header.cell(0, 1), research.coshh_substance_name, bold=True, size=8.1, color=WHITE, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_paragraph("5. Activity-Specific Assessment", style="SectionTitle")
    compact = _compact_activities(assessment)
    table = doc.add_table(rows=len(compact) + 1, cols=6)
    _style_table(table, [1.0, 1.28, 1.55, 2.65, 0.48, 0.58], margin=40)
    for index, value in enumerate(("Area", "Activity", "Exposure", "Controls / PPE", "Initial", "Residual")):
        _shading(table.cell(0, index), NAVY)
        _text(table.cell(0, index), value, bold=True, size=7.0, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    _repeat_header(table.rows[0])
    for row_index, activity in enumerate(compact, 1):
        _no_split(table.rows[row_index])
        controls_and_ppe = _clean_lines(activity.controls[:2] + (["PPE: " + ", ".join(activity.ppe)] if activity.ppe else []), 3)
        row = (
            AREA_LABELS[activity.area],
            activity.activity,
            activity.exposure_scenario,
            controls_and_ppe,
            activity.initial_risk,
            activity.residual_risk,
        )
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            if col_index in (4, 5):
                band = value
                _shading(cell, _risk_fill(band))
                _text(cell, _risk_letter(band), bold=True, size=7.6, align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                _text(cell, str(value), bold=col_index == 0, size=7.0, color=NAVY if col_index == 0 else DARK)

    if assessment.review_flags:
        note = doc.add_paragraph()
        note.paragraph_format.space_before = Pt(2)
        note.paragraph_format.space_after = Pt(2)
        run = note.add_run("Review: ")
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(7.4)
        run.font.color.rgb = RGBColor.from_string(NAVY)
        run = note.add_run(" ".join(assessment.review_flags[:2]))
        run.font.name = "Arial"
        run.font.size = Pt(7.4)

    doc.add_paragraph("6. Emergency, First Aid & Storage", style="SectionTitle")
    emergency = doc.add_table(rows=5, cols=2)
    _style_table(emergency, [1.5, 5.9], margin=40)
    emergency_rows = [
        ("Inhalation", research.first_aid_inhalation),
        ("Skin / eyes", f"Skin: {research.first_aid_skin} Eyes: {research.first_aid_eye}"),
        ("Ingestion", research.first_aid_ingestion),
        ("Spillage", research.spill_response),
        ("Storage / disposal", f"{research.storage_handling} Disposal: {research.disposal}"),
    ]
    for index, (label, value) in enumerate(emergency_rows):
        _shading(emergency.cell(index, 0), LIGHT_TEAL)
        _text(emergency.cell(index, 0), label, bold=True, size=7.4, color=NAVY)
        _text(emergency.cell(index, 1), value, size=7.4)

    doc.add_paragraph("7. SDS Traceability & Human Approval", style="SectionTitle")
    metadata = doc.add_table(rows=3, cols=4)
    _style_table(metadata, [1.1, 2.55, 1.1, 2.65], margin=40)
    metadata_rows = [
        ("SDS", research.sds.manufacturer or "Reference SDS", "SDS status", research.sds.match_status.replace("_", " ").title()),
        ("Revision", research.sds.revision_date or "Not stated", "Original SDS", "Stored unchanged" if stored_sds.status == "STORED_UNCHANGED" else "Source URL retained"),
        ("Health surveillance", "Considered" if research.health_surveillance_considered else "Review", "Recommendation", "Human review required" if research.health_surveillance_recommended else "No specific surveillance identified"),
    ]
    for row_index, row in enumerate(metadata_rows):
        for col_index, value in enumerate(row):
            if col_index % 2 == 0:
                _shading(metadata.cell(row_index, col_index), LIGHT_TEAL)
                _text(metadata.cell(row_index, col_index), value, bold=True, size=7.3, color=NAVY)
            else:
                _text(metadata.cell(row_index, col_index), value, size=7.3)

    approval = doc.add_table(rows=2, cols=4)
    _style_table(approval, [1.05, 2.65, 1.15, 2.55], margin=40)
    approval_rows = [
        ("Reviewed by", "", "Signature / Date", ""),
        ("Approved by", "", "Signature / Date", ""),
    ]
    for row_index, row in enumerate(approval_rows):
        for col_index, value in enumerate(row):
            if col_index % 2 == 0:
                _shading(approval.cell(row_index, col_index), "EAF3F5")
                _text(approval.cell(row_index, col_index), value, bold=True, size=7.4, color=NAVY)
            else:
                _text(approval.cell(row_index, col_index), "", size=7.4)

    for doc_section in doc.sections:
        footer = doc_section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("DRAFT - HUMAN REVIEW AND SIGNATURE REQUIRED")
        run.font.name = "Arial"
        run.font.size = Pt(7.0)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(RED)

    doc.core_properties.title = f"COSHH Assessment - {research.coshh_substance_name}"
    doc.core_properties.subject = "Agent-generated COSHH draft for human review"
    doc.core_properties.author = "Eaststone Material Assessments Agent"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def coshh_review_summary(assessment: COSHHAssessment, stored_sds: StoredSDS) -> str:
    research = assessment.research
    lines = [
        f"COSHH substance: {research.coshh_substance_name}",
        f"Material/use context: {research.material_use_context}",
        f"SDS: {research.sds.manufacturer} - {research.sds.revision_date or 'revision not stated'}",
        f"SDS match: {research.sds.match_status}",
        f"Original SDS storage: {stored_sds.status}",
        f"GHS pictograms: {', '.join(code.value for code in research.pictograms) or 'None identified'}",
        f"Signal word: {research.signal_word}",
        "",
        "Human review flags:",
    ]
    lines.extend(f"- {flag}" for flag in assessment.review_flags) if assessment.review_flags else lines.append("- None generated.")
    lines.extend(
        [
            "",
            "Approval boundary:",
            "This document is a draft. Reviewer/approver names, signatures and dates are intentionally left blank for human completion.",
        ]
    )
    return "\n".join(lines)
